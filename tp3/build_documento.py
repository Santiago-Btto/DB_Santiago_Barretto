"""Genera el informe Word del TP3 a partir de los scripts versionados."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "TP3_Entrega_Santiago_Barretto.docx"
NAVY, PALE, GRID = "17365D", "EAF2F8", "D9D9D9"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)
        borders.append(node)
    props.append(borders)


def cell_format(cell, header=False, alternate=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)
    if header:
        shade(cell, NAVY)
    elif alternate:
        shade(cell, PALE)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            run.font.size = Pt(8.5)
            if header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for index, value in enumerate(headers):
        cell = t.rows[0].cells[index]
        cell.text = value
        cell.width = Inches(widths[index])
        cell_format(cell, header=True)
    for row_index, values in enumerate(rows):
        cells = t.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].width = Inches(widths[index])
            cell_format(cells[index], alternate=row_index % 2 == 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(10.5)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Aptos Display" if level == 1 else "Aptos"
    r.font.color.rgb = RGBColor(0, 0, 0)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)


def make():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Aptos", Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    doc.styles["Title"].font.name = "Aptos Display"
    doc.styles["Title"].font.size = Pt(24)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Trabajo Practico 3")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Optimizacion de consultas sobre Food Store")
    run.bold, run.font.size = True, Pt(15)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Base de Datos II - Unidad 2 - Semana 3\n").bold = True
    meta.add_run("Estudiante: Santiago Barretto\nRepositorio: DB_Santiago_Barretto")
    doc.add_paragraph()
    body(doc, "Este informe resuelve y organiza el TP3: carga masiva de laboratorio, hipotesis de optimizacion, consultas con especificacion precisa y declaracion de uso de IA. Los valores de EXPLAIN ANALYZE se completan exclusivamente con las salidas reales de PostgreSQL.")
    table(doc, ["Elemento", "Estado"], [
        ["Scripts y documentacion", "Preparados en tp3/"],
        ["Planes y tiempos", "Ejecutados en food_store_tp3_mediciones"],
        ["Equivalencias de Parte 4", "Cuatro controles EXCEPT ejecutados con 0 filas distintas"],
        ["DUIA", "Registro completado con decisiones basadas en evidencia real"],
    ], [2.3, 4.1])
    doc.add_page_break()

    heading(doc, "1 Objetivo y protocolo de trabajo")
    body(doc, "El TP3 optimiza consultas reales sobre Food Store con volumen suficiente para medir cambios en el plan. La IA propone scripts, reescrituras e indices; cada decision se acepta solo despues de leerla y contrastarla contra EXPLAIN ANALYZE.")
    heading(doc, "Protocolo de seguridad", 2)
    table(doc, ["Paso", "Aplicacion"], [
        ["Copia", "Crear y usar food_store_tp3_lab; nunca ejecutar sobre la base original."],
        ["Respaldo", "Ejecutar pg_dump sobre la copia antes de cambios estructurales o carga masiva."],
        ["Transaccion", "La carga y los indices tienen controles de error; revisar la salida antes de continuar."],
        ["Evidencia", "Conservar las salidas de tp3/evidencia/ y transcribir solo resultados demostrados."],
    ], [1.25, 5.15])
    heading(doc, "Ejecucion reproducible", 2)
    body(doc, "Desde la carpeta tp3 se ejecuta el siguiente comando. Genera preflight, carga, planes antes y despues, indices, equivalencias y validacion.")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.2)
    r = code.add_run(".\\run_mediciones.ps1 -Database food_store_tp3_lab -Lote SB2026TP3 -PrepararDatos")
    r.font.name, r.font.size = "Consolas", Pt(8.5)

    heading(doc, "2 Parte 1 Carga masiva")
    body(doc, "01_carga_masiva.sql usa generate_series dentro de una transaccion y no elimina ni modifica filas existentes. Al finalizar, ANALYZE actualiza estadisticas de cliente, producto, pedido y detalle_pedido.")
    table(doc, ["Entidad", "Cantidad", "Controles"], [
        ["Producto", "50.000", "Categorias existentes, precio 500 a 5000, stock 0 a 200."],
        ["Cliente", "20.000", "Email unico e identificacion por lote."],
        ["Pedido", "200.000", "Cliente existente, fecha generada y forma de pago valida."],
        ["Detalle pedido", "400.000", "Dos productos distintos por pedido, PK compuesta y valores positivos."],
    ], [1.1, 1.0, 4.3])
    body(doc, "La validacion posterior exige estos conteos, los rangos de producto, la unicidad de email y la ausencia de referencias huerfanas.")
    doc.add_page_break()

    heading(doc, "3 Parte 2 Consultas lentas y optimizacion medida")
    body(doc, "Las mismas tres consultas se ejecutan antes y despues. Los indices son hipotesis de mejora: la decision final depende del nodo elegido y del tiempo real del plan.")
    table(doc, ["Consulta", "Patron", "Cambio propuesto"], [
        ["Q1", "Productos activos por categoria, rango de precio, orden y LIMIT.", "Indice parcial: categoria, precio DESC e identificador para activos."],
        ["Q2", "Pedidos recientes por forma de pago, fecha DESC y LIMIT.", "Indice por forma de pago y fecha DESC."],
        ["Q3", "Productos mas vendidos de los ultimos siete dias.", "Indice por fecha de pedido para acotar el rango antes de los joins."],
    ], [0.55, 2.45, 3.4])
    heading(doc, "Tabla de resultados", 2)
    table(doc, ["Consulta", "Plan antes", "Cambio", "Plan despues", "Mejora"], [
        ["Q1", "Bitmap Heap Scan y Sort; 44.727 ms", "Indice parcial de producto", "Bitmap Index Scan; 45.011 ms", "0.994x: no mejoro"],
        ["Q2", "Parallel Seq Scan y Sort; 26.940 ms", "Indice forma_pago y fecha", "Index Scan; 0.553 ms", "48.72x"],
        ["Q3", "Parallel Seq Scan de pedido; 68.061 ms", "Indice fecha de pedido", "Bitmap Index Scan; 52.491 ms", "1.30x"],
    ], [0.45, 1.75, 1.15, 1.75, 0.7])
    body(doc, "Mejora x = tiempo total antes / tiempo total despues. cost es una estimacion del planificador; no representa milisegundos. Q1 no se acepta como mejora: aunque el indice se uso, la consulta que calcula la categoria siguio haciendo un Seq Scan y el tiempo total fue ligeramente mayor.")

    heading(doc, "4 Parte 3 Lectura critica de un plan")
    body(doc, "Elegir un plan posterior real, pegarlo junto con la respuesta literal de IA y contrastar cada afirmacion. Distinguir cost, actual time, rows, loops, buffers y Rows Removed by Filter.")
    table(doc, ["Afirmacion de IA", "Correcta", "Evidencia del plan real"], [
        ["El indice por forma de pago y fecha evita leer todos los pedidos y evita un Sort final.", "Si", "Index Scan idx_tp3_pedido_forma_pago_fecha; no hay nodo Sort; Index Cond contiene forma_pago y fecha."],
        ["El costo final 143.13 representa que la consulta tardo 143.13 milisegundos.", "No", "cost=0.42..143.13 es estimado. El tiempo real es Execution Time: 0.553 ms."],
        ["Los 200 resultados se obtuvieron con pocas lecturas de buffers.", "Si", "Buffers: shared hit=200 read=4 y 200 filas devueltas por Index Scan."],
    ], [2.3, 0.8, 2.7])
    doc.add_page_break()

    heading(doc, "5 Parte 4 Consultas con especificacion precisa")
    heading(doc, "Spec 1 Resumen por categoria", 2)
    body(doc, "Devuelve, para cada categoria activa, identificador, nombre, cantidad de productos activos y valor de lista total. Incluye categorias sin productos con cero; ordena por cantidad descendente, nombre e identificador. La version IA usa LEFT JOIN y agregacion; la alternativa propia usa subconsultas correlacionadas.")
    heading(doc, "Spec 2 Productos sin ventas recientes", 2)
    body(doc, "Devuelve productos y categorias activos que no tuvieron detalles asociados a pedidos de los ultimos 90 dias. La version IA usa NOT IN y la alternativa propia NOT EXISTS. detalle_pedido.id_producto no admite NULL por su clave primaria compuesta.")
    heading(doc, "Equivalencia formal", 2)
    table(doc, ["Control", "Esperado", "Resultado real"], [
        ["Spec 1 IA menos propia", "0 filas", "0 filas distintas"],
        ["Spec 1 propia menos IA", "0 filas", "0 filas distintas"],
        ["Spec 2 IA menos propia", "0 filas", "0 filas distintas"],
        ["Spec 2 propia menos IA", "0 filas", "0 filas distintas"],
    ], [2.3, 1.15, 2.35])

    heading(doc, "6 Parte 5 Competencia de optimizacion")
    body(doc, "La consulta comun lista productos y categorias activos con precio entre 1500 y 4500, ordenados por precio descendente e identificador, con LIMIT 100. Se mide antes de la estrategia y luego se repite en la misma copia.")
    table(doc, ["Equipo", "Estrategia", "Antes ms", "Despues ms", "Mejora x"], [[
        "Santiago Barretto", "Indice parcial precio DESC e identificador con INCLUDE", "22.246", "0.549", "40.52x"
    ]], [1.2, 2.7, 0.9, 1.05, 0.8])
    doc.add_page_break()

    heading(doc, "7 Declaracion de Uso de IA")
    body(doc, "La IA se utilizo para proponer scripts y documentacion. Se revisa cada propuesta como hipotesis y la aceptacion final depende de PostgreSQL y de la evidencia conservada.")
    table(doc, ["Uso", "Especificacion resumida", "Decision y verificacion"], [
        ["Carga masiva", "Generar 50k productos, 20k clientes, 200k pedidos y detalles con PK, FK, UNIQUE y CHECK.", "Aceptada y ejecutada: conteos reales validados sin referencias huerfanas."],
        ["Q1 Q2 Q3", "Proponer consultas medibles y un indice por patron, sin afirmar mejora.", "Q2 y Q3 confirmados; Q1 documentado como no mejorado."],
        ["Parte 4", "Una consulta resumen y una subconsulta, con filtros activos y EXCEPT bidireccional.", "Aceptadas: los cuatro controles devolvieron cero."],
        ["Parte 3 y 5", "Explicar plan real y proponer estrategia para consulta comun.", "Parte 3 contrastada contra Q2; competencia confirmada con mejora 40.52x."],
    ], [1.1, 2.85, 2.65])
    heading(doc, "8 Checklist antes de entregar")
    for text in [
        "Ejecute el runner sobre una copia de laboratorio.",
        "Guarde planes antes, indices y planes despues en evidencia/.",
        "Complete la tabla de resultados con nodos y tiempos reales, no estimados.",
        "Pegue un plan posterior y la explicacion literal de IA en la Parte 3.",
        "Confirme los cuatro EXCEPT de Parte 4 con resultado cero.",
        "Registre propuestas aceptadas y descartadas en la DUIA.",
        "Puedo explicar cada script e indice incluido en la entrega.",
    ]:
        bullet(doc, text)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Base de Datos II - TP3 - Food Store").font.size = Pt(8)
    doc.core_properties.title = "Trabajo Practico 3 Food Store"
    doc.core_properties.author = "Santiago Barretto"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make()
